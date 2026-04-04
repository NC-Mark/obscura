"""
Text and DOCX anonymization using PIIEngine.
"""

from __future__ import annotations

import json
import logging
import time
import uuid
from collections import defaultdict
from pathlib import Path

from ..storage import save_mapping, save_review
from .docx_utils import (
    iter_docx_paragraphs, iter_all_wp_elements,
    get_runs, replace_in_runs,
    replace_across_runs, replace_cross_paragraphs,
    save_docx, docx_to_html,
)

log = logging.getLogger("pii-shield.anonymizer")
_flog = logging.getLogger("pii-shield.ner")


def anonymize_text(engine, text: str, language: str = "en",
                   prefix: str = "", entity_overrides: str = "",
                   profile=None) -> dict:
    """
    Detect and anonymize PII in plain text.

    Returns a dict with:
      anonymized_text, session_id, total_entities, entities_confirmed,
      unique_entities, by_type, entities (safe — no real values), processing_time_ms
    """
    t0 = time.time()
    _flog.info(f"=== anonymize_text | len={len(text)} | prefix={prefix!r} ===")

    entities = engine.detect(text, language, profile=profile)
    confirmed = [e for e in entities if e.get("verified")]

    if entity_overrides:
        confirmed = engine._apply_overrides(confirmed, text, entity_overrides)

    mapping = engine._assign_placeholders(confirmed, prefix)

    # Apply substitutions right-to-left to preserve offsets
    anonymized = text
    for e in sorted(confirmed, key=lambda x: x["start"], reverse=True):
        anonymized = anonymized[:e["start"]] + e["placeholder"] + anonymized[e["end"]:]

    session_id = uuid.uuid4().hex[:12]
    save_mapping(session_id, mapping, {"confirmed": len(confirmed)})

    by_type: dict = defaultdict(int)
    for e in confirmed:
        by_type[e["type"]] += 1

    safe_entities = [
        {
            "placeholder": e.get("placeholder", ""),
            "type": e["type"],
            "score": e["score"],
            "verified": e["verified"],
            "reason": e.get("reason", ""),
        }
        for e in confirmed
    ]

    # Store review data for optional HITL step
    review_data = {
        "original_text": text,
        "domain": getattr(engine, "_domain", "general"),
        "entities": [
            {
                "type": e["type"], "text": e["text"],
                "start": e["start"], "end": e["end"],
                "score": e["score"], "verified": e.get("verified", False),
            }
            for e in entities
        ],
        "confirmed": [i for i, e in enumerate(entities) if e.get("verified")],
        "status": "pending",
        "overrides": {"remove": [], "add": []},
        "timestamp": time.time(),
    }
    save_review(session_id, review_data)

    return {
        "anonymized_text": anonymized,
        "session_id": session_id,
        "total_entities": len(entities),
        "entities_confirmed": len(confirmed),
        "unique_entities": len(mapping),
        "by_type": dict(by_type),
        "entities": safe_entities,
        "processing_time_ms": round((time.time() - t0) * 1000, 1),
    }


def anonymize_docx(engine, docx_path: str | Path,
                   language: str = "en", prefix: str = "") -> dict:
    """
    Anonymize PII in a .docx file, preserving formatting.
    Detects entity-by-entity per paragraph via python-docx run API.
    """
    from docx import Document
    t0 = time.time()
    docx_path = Path(docx_path)
    doc = Document(str(docx_path))

    type_counters: dict = defaultdict(int)
    seen_exact: dict = {}
    seen_family: dict = {}
    mapping: dict = {}
    total = 0
    by_type: dict = defaultdict(int)

    for para in iter_docx_paragraphs(doc):
        full_text, runs_info = get_runs(para)
        if not full_text.strip():
            continue

        entities = engine.detect(full_text, language)
        confirmed = [e for e in entities if e.get("verified")]

        for e in sorted(confirmed, key=lambda x: x["start"], reverse=True):
            ph = engine._get_or_create_placeholder(
                e["type"], e["text"], type_counters,
                seen_exact, seen_family, mapping, prefix
            )
            replace_in_runs(runs_info, e["start"], e["end"], ph)
            total += 1
            by_type[e["type"]] += 1

    session_id = uuid.uuid4().hex[:12]
    out_dir = docx_path.parent / f"pii_shield_{session_id}"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{docx_path.stem}_anonymized.docx"
    save_docx(doc, out_path)
    save_mapping(session_id, mapping, {"source": str(docx_path)})

    return {
        "output_path": str(out_path),
        "output_dir": str(out_dir),
        "session_id": session_id,
        "total_entities": total,
        "unique_entities": len(mapping),
        "by_type": dict(by_type),
        "processing_time_ms": round((time.time() - t0) * 1000, 1),
    }


def anonymize_docx_with_mapping(engine, docx_path: str | Path,
                                 mapping: dict, out_dir: str | Path = None) -> str:
    """
    Apply an existing placeholder mapping to a .docx via find-replace.
    Used for re-anonymization after HITL review (no new NER detection needed).
    """
    from docx import Document
    docx_path = Path(docx_path)
    doc = Document(str(docx_path))
    wns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

    reverse_map = {v: k for k, v in mapping.items()}
    sorted_texts = sorted(reverse_map.keys(), key=len, reverse=True)

    all_p_elems = list(iter_all_wp_elements(doc))
    _flog.info(f"anonymize_docx_with_mapping | {len(mapping)} placeholders | "
               f"{len(all_p_elems)} paragraphs")

    cross_para_texts = []
    for real_text in sorted_texts:
        found = False
        for p_elem in all_p_elems:
            from .docx_utils import collect_paragraph_segments
            segs = collect_paragraph_segments(p_elem, wns)
            vtext = "".join(s[1] for s in segs)
            if real_text in vtext:
                replace_across_runs(p_elem, real_text, reverse_map[real_text], wns)
                found = True
        if not found and "\n" in real_text:
            cross_para_texts.append(real_text)

    for real_text in cross_para_texts:
        all_p_fresh = list(iter_all_wp_elements(doc))
        replace_cross_paragraphs(all_p_fresh, real_text, reverse_map[real_text], wns)

    parent = Path(out_dir) if out_dir else docx_path.parent
    out = parent / f"{docx_path.stem}_anonymized.docx"
    save_docx(doc, out)
    return str(out)


def anonymize_file(engine, file_path: str | Path, language: str = "en",
                   prefix: str = "", review_session_id: str = "",
                   profile=None) -> dict:
    """
    Unified entry point: auto-detects format and anonymizes.
    Supports .pdf, .docx, .txt, .md, .csv, .log

    If review_session_id is provided, loads HITL overrides from disk
    and re-anonymizes with corrections applied.
    """
    from ..storage import get_review, load_mapping as _load_mapping

    p = Path(file_path).expanduser().resolve()
    if not p.exists():
        return {"error": f"File not found: {p}"}

    entity_overrides = ""
    if review_session_id:
        review = get_review(review_session_id.strip())
        if review:
            overrides = review.get("overrides", {})
            if overrides.get("remove") or overrides.get("add"):
                import json as _json
                entity_overrides = _json.dumps(overrides)
        else:
            return {"error": f"Review session not found: {review_session_id}"}

    suffix = p.suffix.lower()

    def _make_out_dir(session_id: str) -> Path:
        d = p.parent / f"pii_shield_{session_id}"
        d.mkdir(parents=True, exist_ok=True)
        return d

    if suffix == ".pdf":
        try:
            import pdfplumber
        except ImportError:
            return {"error": "pdfplumber not installed. Run: pip install pdfplumber"}
        try:
            with pdfplumber.open(str(p)) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        except Exception as e:
            return {"error": f"Failed to read PDF: {e}"}
        if len(text.strip()) < 50:
            return {"error": "PDF has no extractable text (scanned/image PDFs not supported)."}
        r = anonymize_text(engine, text, language, prefix=prefix,
                           entity_overrides=entity_overrides)
        out_dir = _make_out_dir(r["session_id"])
        out = out_dir / f"{p.stem}_anonymized.txt"
        out.write_text(r["anonymized_text"], encoding="utf-8")
        r.pop("anonymized_text", None)
        r["output_path"] = str(out)
        r["output_dir"] = str(out_dir)
        return r

    elif suffix == ".docx":
        from docx import Document as _DocxDoc
        try:
            _doc = _DocxDoc(str(p))
            text = "\n".join(para.text for para in iter_docx_paragraphs(_doc))
        except Exception as e:
            return {"error": f"Failed to read .docx: {e}"}

        # Generate HTML for review UI
        docx_html = None
        try:
            docx_html = docx_to_html(_doc)
        except Exception:
            pass

        r = anonymize_text(engine, text, language, prefix=prefix,
                           entity_overrides=entity_overrides)
        out_dir = _make_out_dir(r["session_id"])

        # Attach HTML to review data
        if docx_html:
            from ..storage import get_review, save_review
            rv = get_review(r["session_id"])
            if rv:
                rv["original_html"] = docx_html
                save_review(r["session_id"], rv)

        out_txt = out_dir / f"{p.stem}_anonymized.txt"
        out_txt.write_text(r["anonymized_text"], encoding="utf-8")
        r.pop("anonymized_text", None)
        r["output_path"] = str(out_txt)
        r["output_dir"] = str(out_dir)

        # Also produce formatted anonymized .docx
        try:
            mapping = _load_mapping(r["session_id"])
            docx_out = anonymize_docx_with_mapping(engine, p, mapping, out_dir)
            r["docx_output_path"] = docx_out
        except Exception as e:
            log.warning(f"anonymize_docx_with_mapping failed (txt OK): {e}")

        return r

    elif suffix in (".txt", ".md", ".csv", ".log", ".text"):
        text = p.read_text(encoding="utf-8")
        r = anonymize_text(engine, text, language, prefix=prefix,
                           entity_overrides=entity_overrides)
        out_dir = _make_out_dir(r["session_id"])
        out = out_dir / f"{p.stem}_anonymized{suffix}"
        out.write_text(r["anonymized_text"], encoding="utf-8")
        r.pop("anonymized_text", None)
        r["output_path"] = str(out)
        r["output_dir"] = str(out_dir)
        return r

    else:
        return {"error": f"Unsupported format: {suffix}. "
                         f"Supported: .pdf .docx .txt .md .csv .log"}
